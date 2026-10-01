#!/usr/bin/env python3
"""Bounded research then deterministic review-only Google Docs delivery."""
from __future__ import annotations

import argparse
import hashlib
import ipaddress
import json
import os
import re
import socket
import subprocess
import sys
import uuid
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from artifact_registry import record_artifact
from google_doc_delivery import atomic_json, deliver_document

WORKSPACE = Path('/var/lib/openclaw/.openclaw/workspace')
MODEL = 'openai-codex/gpt-5.6-sol'
JOB = 'xhs-recommendation-every-3-days'


def configured_recipient() -> str:
    # Identity is private host configuration, never a model prompt or CLI argument.
    value = os.environ.get('OPENCLAW_OWNER_GOOGLE_EMAIL', '')
    path = Path('/etc/openclaw/openclaw.env')
    if not value and path.is_file():
        for line in path.read_text().splitlines():
            if line.startswith('OPENCLAW_OWNER_GOOGLE_EMAIL='):
                value = line.split('=', 1)[1].strip().strip('\"').strip("'")
    if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', value):
        raise ValueError('protected Google recipient configuration is missing')
    return value


def validate_public_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password or parsed.port not in (None, 443):
        raise ValueError('image/source URL must be public HTTPS')
    addresses = socket.getaddrinfo(parsed.hostname, 443, type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
        raise ValueError('non-public image/source address is blocked')
    return url


def validate_manifest(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError('manifest must be an object')
    for key in ('product', 'title', 'body', 'version'):
        if not isinstance(payload.get(key), str) or not payload[key].strip():
            raise ValueError(f'missing manifest field: {key}')
    if len(payload['title']) > 40 or '日本 Costco' in payload['title']:
        raise ValueError('title violates recommendation rules')
    if len(payload['body']) > 4000 or 'コストコ' not in payload['body']:
        raise ValueError('body length or Japan context is invalid')
    tags = payload.get('tags')
    if not isinstance(tags, list) or not tags or any(not isinstance(t,str) or not t.startswith('#') or '\n' in t for t in tags):
        raise ValueError('tags must be individual hashtag lines')
    images = payload.get('images')
    if not isinstance(images, list) or len(images) != 3:
        raise ValueError('exactly three verified images are required')
    if sum(i.get('kind') == 'official' for i in images) != 2 or sum(i.get('kind') == 'real_photo' for i in images) != 1:
        raise ValueError('two official images and one real photo are required')
    for image in images:
        if image.get('version') != payload['version'] or image.get('no_watermark') is not True or image.get('real') is not True:
            raise ValueError('image version/authenticity verification is missing')
        validate_public_url(str(image.get('url') or ''))
        validate_public_url(str(image.get('source') or ''))
    if len({i['url'] for i in images}) != 3:
        raise ValueError('duplicate image URLs')
    return payload


def fetch_image(url: str, destination: Path) -> str:
    import requests
    from PIL import Image, ImageOps
    current = validate_public_url(url)
    for redirect in range(4):
        with requests.get(current, timeout=(15, 45), stream=True, allow_redirects=False) as response:
            if response.is_redirect:
                from urllib.parse import urljoin
                current = validate_public_url(urljoin(current, response.headers['Location']))
                continue
            response.raise_for_status()
            if not response.headers.get('Content-Type','').lower().startswith('image/'):
                raise ValueError('image response has unexpected MIME type')
            content = bytearray()
            for chunk in response.iter_content(65536):
                content.extend(chunk)
                if len(content) > 10_000_000:
                    raise ValueError('image exceeds size limit')
            break
    else:
        raise ValueError('too many image redirects')
    import io
    with Image.open(io.BytesIO(content)) as image:
        image.load()
        if min(image.size) < 200 or image.width * image.height > 30_000_000:
            raise ValueError('image dimensions are unsuitable')
        ImageOps.exif_transpose(image).convert('RGB').save(destination, format='JPEG', quality=95)
    return hashlib.sha256(destination.read_bytes()).hexdigest()


def build_docx(manifest: dict, directory: Path) -> Path:
    from docx import Document
    from docx.shared import Inches
    validate_manifest(manifest)
    document = Document()
    document.add_heading(manifest['title'], 0)
    for paragraph in manifest['body'].split('\n'):
        if paragraph.strip():
            document.add_paragraph(paragraph)
    for tag in manifest['tags']:
        document.add_paragraph(tag)
    hashes = set()
    for index, image in enumerate(manifest['images'], 1):
        destination = directory / f'image-{index}.jpg'
        digest = fetch_image(image['url'], destination)
        if digest in hashes:
            raise ValueError('image content is duplicated')
        hashes.add(digest)
        document.add_picture(str(destination), width=Inches(5.5))
    document.add_heading('配图与核验来源（供确认）', 1)
    document.add_paragraph(f"商品：{manifest['product']}；版本：{manifest['version']}")
    for image in manifest['images']:
        document.add_paragraph(image['source'])
    path = directory / 'draft.docx'
    document.save(path)
    return path


def prepare_manifest(directory: Path) -> Path:
    manifest = directory / 'manifest.json'
    if manifest.is_file():
        return manifest
    rules = (WORKSPACE / 'XHS_RECOMMENDATION_RULES.md').read_text(encoding='utf-8')
    rules = re.sub(r'[\w.+-]+@[\w.-]+', '[private]', rules)
    prompt = (
        'Prepare one review-only Xiaohongshu recommendation. Do not publish or use Google Docs. '
        'No browser tool is available. Use web_search/web_fetch for current product and image evidence. '
        'At most 8 source fetches, short extracts only, at most one retry per failed source. '
        'Never fabricate personal ownership or use experience. Stop with a failure if evidence is insufficient. '
        f'Follow these rules:\n{rules[:6500]}\n'
        f'Write UTF-8 JSON to {manifest}. Return only the path and a brief outcome. '
        'Schema: product:string,title:string,body:string,version:string,tags:["#tag"],'
        'images:[{kind:"official"|"real_photo",url:https_image_url,source:https_page_url,'
        'version:exact_same_version,real:true,no_watermark:true}]. '
        'Exactly two official and one Japanese real-photo image, distinct compositions, same packaging version. '
        'Title in Chinese, body uses コストコ, tags one per item. '
        'Retain evidence notes in a separate sources.md alongside the manifest. '
        'Do not fetch full-page HTML, run browser/CDP code, read unrelated files or install packages.'
    )
    env = dict(os.environ, HOME='/var/lib/openclaw')
    env.pop('OPENCLAW_OWNER_GOOGLE_EMAIL', None)
    result = subprocess.run(
        ['openclaw','--no-color','agent','--agent','xhs-writer','--session-id',str(uuid.uuid4()),
         '--message',prompt,'--timeout','1200','--thinking','low','--json'],
        env=env, capture_output=True, text=True, timeout=1260)
    if result.returncode or not manifest.is_file() or manifest.stat().st_size > 50000:
        raise ValueError('bounded product research did not produce a manifest')
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--run-date', default=datetime.now(ZoneInfo('Asia/Tokyo')).date().isoformat())
    parser.add_argument('--prepared-docx', type=Path)
    args = parser.parse_args()
    try:
        run_date = datetime.strptime(args.run_date, '%Y-%m-%d').date().isoformat()
        directory = WORKSPACE / 'state' / 'xhs-delivery' / run_date
        directory.mkdir(parents=True, exist_ok=True)
        import fcntl
        with (directory.parent / '.lock').open('w') as lock:
            try:
                fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError:
                raise ValueError('another XHS delivery is running')
            recipient = configured_recipient()
            draft = args.prepared_docx or directory / 'draft.docx'
            if not draft.is_file():
                payload = json.loads(prepare_manifest(directory).read_text(encoding='utf-8'))
                draft = build_docx(payload, directory)
            receipt = deliver_document(draft, recipient, directory / 'receipt.json')
            record_artifact(receipt['url'], JOB)
            print(f"小红书推荐文草稿已完成。\n正文、三张图片和指定账号的查看权限已核验。\n{receipt['url']}\n等待你确认，尚未发布小红书。")
            return 0
    except Exception as exc:
        # Playwright errors can contain account text; only a non-sensitive class is logged.
        print(f'XHS delivery failed: {type(exc).__name__}; no successful delivery receipt.', file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
